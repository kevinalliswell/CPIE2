#!/usr/bin/env python3
"""Exercise the real application in an isolated, hardware-free subprocess.

    python scripts/smoke_check.py --source --output artifacts/source-smoke.json
    python scripts/smoke_check.py --executable dist/CPIE/CPIE.exe --output artifacts/frozen-smoke.json

The explicit app --smoke-test mode uses a fresh temporary workspace. It never
opens the operator's database, modifies camera calibration, or connects devices.
"""
import argparse
import json
import logging
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import traceback


def _write_report(path, report):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def _offline_connection_shutdown(app, window_class, report):
    """Exercise failed discovery and real window close with only fake transports."""
    from PySide6.QtCore import QCoreApplication, QEvent
    from PySide6.QtWidgets import QMessageBox
    from pymodbus.exceptions import ModbusIOException
    import modbus_multi_device.manager as manager_module

    transports = []

    class OfflineTransport:
        def __init__(self, **kwargs):
            self.closed = False
            self.reads = 0
            self.writes = 0
            transports.append(self)

        def connect(self):
            return True  # Only the adapter opens; no field device ever responds.

        def close(self):
            self.closed = True

        def no_response(self, *args, **kwargs):
            self.reads += 1
            raise ModbusIOException('Synthetic offline device: no response')

        read_holding_registers = no_response
        read_input_registers = no_response
        read_coils = no_response

        def reject_write(self, *args, **kwargs):
            self.writes += 1
            raise AssertionError('Untouched offline discovery must not send control commands')

        write_coil = reject_write
        write_register = reject_write
        write_registers = reject_write

    original_client = manager_module.ModbusSerialClient
    original_critical = QMessageBox.critical
    offline_window = None
    controllers = []
    connection_dialogs = []
    try:
        manager_module.ModbusSerialClient = OfflineTransport
        offline_window = window_class()
        pages = [offline_window.stacked_widget.widget(index) for index in (1, 2)]
        controllers = [page.controller for page in pages]

        def expected_connection_failure(parent, title, message, *args, **kwargs):
            if parent in pages and title == '连接错误':
                connection_dialogs.append(message)
                return QMessageBox.Ok
            return original_critical(parent, title, message, *args, **kwargs)

        QMessageBox.critical = staticmethod(expected_connection_failure)
        offline_window.show()
        for page in pages:
            page._on_connect()
        for controller in controllers:
            assert controller._connect_thread is not None
            controller._connect_thread.join(5.0)
            assert not controller._connect_thread.is_alive(), 'Offline discovery did not finish'
        app.processEvents()
        assert len(connection_dialogs) == 2, connection_dialogs
        assert all(controller.current_session_id is None and controller.current_state.name == 'IDLE'
                   for controller in controllers)
        assert all(controller.manager.connected and not controller.manager.ready
                   and not controller.manager.started for controller in controllers)
        assert len(transports) == 2 and all(transport.reads > 0 for transport in transports)
        assert offline_window.close(), 'Application refused exit after entirely offline discovery'
        assert all(controller.db is None and controller.manager is None for controller in controllers)
        assert all(transport.closed and transport.writes == 0 for transport in transports)
        report['offline_connection_shutdown'] = True
    finally:
        QMessageBox.critical = staticmethod(original_critical)
        manager_module.ModbusSerialClient = original_client
        # On a regression, stop only workers using these fake transports before
        # deleting the test window. This never reaches a real serial device.
        for controller in controllers:
            if controller.manager is not None and controller.manager.client in transports:
                controller.manager.disconnect()
        if offline_window is not None:
            offline_window.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


def application_smoke(output):
    """Called only by the executable's explicit --smoke-test entry point."""
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    os.environ.setdefault('MPLBACKEND', 'Agg')
    report = {'success': False, 'frozen': bool(getattr(sys, 'frozen', False)), 'hardware_access': []}
    output = Path(output).resolve()
    asset_root = Path(sys.executable).parent if report['frozen'] else Path(__file__).resolve().parents[1]
    original_cwd = Path.cwd()
    workspace = tempfile.TemporaryDirectory(prefix='cpie-smoke-')
    root = Path(workspace.name)
    report['workspace'] = str(root)
    lock = None
    app = None
    window = None
    try:
        for name in ('configs', 'resources'):
            shutil.copytree(asset_root / name, root / name)
        os.chdir(root)
        from src.utils.path_manager import PathManager as QualifiedPaths
        from utils.path_manager import PathManager as LegacyPaths
        for paths in (QualifiedPaths, LegacyPaths):
            paths.get_project_root = staticmethod(lambda: str(root))

        # Prevent even an accidental connection introduced by a future startup
        # change. These are boundary guards, not replacement device simulators.
        def reject_hardware(*args, **kwargs):
            report['hardware_access'].append('unexpected device access')
            raise RuntimeError('Smoke mode forbids hardware access')

        import serial
        from pymodbus.client import ModbusSerialClient
        import flamekit
        from flamekit.camera import CameraCapture
        serial.Serial.open = reject_hardware
        ModbusSerialClient.connect = reject_hardware
        CameraCapture.initialize = reject_hardware
        camera_defaults = Path(flamekit.__file__).with_name('config_default.json')
        shutil.copy2(camera_defaults, root / 'config.json')
        report['flamekit_defaults'] = True

        from PySide6.QtWidgets import QApplication, QMessageBox
        from PySide6.QtCore import QTimer, QCoreApplication, QEvent
        from src.views.main_window import MainWindow
        from src.utils.single_instance import SingleInstance
        from src.utils.tools import Tools
        from src.app import recover_interrupted_sessions
        from models.explosion_database import ExplosionDatabase
        from models.ignition_database import IgnitionDatabase

        app = QApplication([])
        app.setQuitOnLastWindowClosed(False)
        lock = SingleInstance('CPIE-smoke-' + str(os.getpid()))
        if not lock.try_lock():
            raise RuntimeError('Cannot acquire isolated smoke lock')

        def unexpected_dialog(parent, title, message, *args, **kwargs):
            report.setdefault('unexpected_dialogs', []).append({'title': title, 'message': message})
            return QMessageBox.No

        for name in ('warning', 'critical', 'information'):
            setattr(QMessageBox, name, staticmethod(unexpected_dialog))
        QMessageBox.question = staticmethod(lambda *args, **kwargs: QMessageBox.Yes)

        explosion = ExplosionDatabase(QualifiedPaths.get_data_path('explosion_experiment.db'))
        explosion_id = explosion.start_experiment_session(experiment_name='Smoke recovery explosion', sample_name='Smoke coal')
        assert explosion_id > 0
        assert explosion.add_test_round(explosion_id, 1, 30.0, '') > 0
        explosion.close()
        ignition = IgnitionDatabase(QualifiedPaths.get_data_path('ignition_experiment.db'))
        ignition_id = ignition.start_experiment_session(experiment_id='SMOKE-IGN', experiment_name='Smoke recovery ignition', sample_names='["Smoke coal"]')
        assert ignition_id > 0
        assert ignition.insert_ignition_data(
            pv=300.0, ch1=280.0, ch2=281.0, ch3=282.0,
            ch4=283.0, ch5=284.0, ch6=285.0, session_id=ignition_id,
        ) > 0
        ignition.close()

        Tools.apply_stylesheet('dark')
        window = MainWindow()
        file_handlers = [handler for handler in logging.getLogger().handlers
                         if isinstance(handler, logging.FileHandler)]
        assert file_handlers, 'Application file logging was not initialized'
        for handler in file_handlers:
            Path(handler.baseFilename).resolve().relative_to(root.resolve())
        report['logs_isolated'] = True
        recovered = recover_interrupted_sessions(window)
        assert recovered == {'explosion': 1, 'ignition': 1}, recovered
        report['recovered_sessions'] = recovered
        window.show()
        result_code = [1]

        def verify_and_close():
            nonlocal window
            try:
                pages = [type(window.stacked_widget.widget(i)).__name__ for i in range(window.stacked_widget.count())]
                assert len(pages) == 7, pages
                for index in range(len(pages)):
                    window.stacked_widget.setCurrentIndex(index)
                    app.processEvents()
                controllers = [window.stacked_widget.widget(index).controller for index in (1, 2)]
                assert all(controller.manager is None for controller in controllers)
                assert all(controller.current_session_id is None for controller in controllers)
                assert all(controller.current_state.name == 'IDLE' for controller in controllers)
                assert window.secondary_display_window is not None, 'Secondary display did not open'
                report['pages'] = pages
                report['secondary_display'] = True

                history = window.stacked_widget.widget(3)
                history.load_experiments()
                assert len(history.experiments_data) == 2
                assert all(item['status'] == 'error' for item in history.experiments_data)
                export_path = root / 'exports' / 'smoke.xlsx'
                export_path.parent.mkdir(exist_ok=True)
                assert history._export_to_excel(str(export_path), [explosion_id], 'explosion', True, True, include_rounds=True)
                from openpyxl import load_workbook
                workbook = load_workbook(export_path, read_only=True)
                assert workbook.active.max_row == 2
                assert workbook['测试轮次详情'].max_row == 2
                workbook.close()
                csv_path = root / 'exports' / 'smoke.csv'
                assert history._export_to_csv(str(csv_path), [ignition_id], 'ignition', True, True)
                assert 'SMOKE-IGN' in csv_path.read_text(encoding='utf-8-sig')
                from PySide6.QtWidgets import QFileDialog
                from src.views.dialogs.generate_report_dialog import GenerateReportDialog
                from docx import Document
                word_path = root / 'exports' / 'smoke.docx'
                original_save_dialog = QFileDialog.getSaveFileName
                original_information = QMessageBox.information
                QFileDialog.getSaveFileName = staticmethod(lambda *args, **kwargs: (str(word_path), ''))
                QMessageBox.information = staticmethod(lambda *args, **kwargs: QMessageBox.Ok)
                word_dialog = GenerateReportDialog(
                    explosion_id, 'explosion',
                    {'experiment_code': 'SMOKE-EXP', 'sample_name': 'Smoke coal',
                     'tests': [{'test_sequence': 1, 'flame_length_mm': 30.0}],
                     'conclusion': 'Synthetic smoke verification only'},
                    explosion_db=history.explosion_db,
                )
                try:
                    word_dialog.on_export_word()
                    document = Document(word_path)
                    assert 'SMOKE-EXP' in '\n'.join(paragraph.text for paragraph in document.paragraphs)
                finally:
                    word_dialog.close()
                    QFileDialog.getSaveFileName = staticmethod(original_save_dialog)
                    QMessageBox.information = staticmethod(original_information)
                report['history_exports'] = ['xlsx', 'csv', 'docx']

                # Closing follows the same prepare-shutdown/cleanup path as a
                # user-confirmed exit. Any rejected close is a failed smoke.
                assert window.close(), 'Application refused clean shutdown'
                for database_class, name, expected_id in (
                    (ExplosionDatabase, 'explosion', explosion_id),
                    (IgnitionDatabase, 'ignition', ignition_id),
                ):
                    reopened = database_class(QualifiedPaths.get_data_path(name + '_experiment.db'))
                    sessions = reopened.get_all_experiment_sessions()
                    assert any(row['id'] == expected_id and row['status'] == 'error' for row in sessions)
                    assert reopened.recover_interrupted_sessions() == 0
                    reopened.close()
                window.deleteLater()
                window = None
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                _offline_connection_shutdown(app, MainWindow, report)
                assert not report['hardware_access'], report['hardware_access']
                assert not report.get('unexpected_dialogs'), report.get('unexpected_dialogs')
                report['database_reopen'] = True
                report['success'] = True
                result_code[0] = 0
            except BaseException:
                report['error'] = traceback.format_exc()
            finally:
                app.exit(result_code[0])

        # Schedule after MainWindow construction so its delayed secondary
        # display is exercised even on a slow first Qt launch.
        QTimer.singleShot(1800, verify_and_close)
        app.exec()
        return result_code[0]
    except BaseException:
        report['error'] = traceback.format_exc()
        return 1
    finally:
        if window is not None and window.isVisible():
            window.close()
        if lock is not None:
            lock.unlock()
        _write_report(output, report)
        os.chdir(original_cwd)
        # Windows can retain Qt/SQLite/log handles until process exit. The
        # external runner removes its own artifacts; retained smoke temp files
        # contain synthetic data only and never interfere with operator files.
        try:
            workspace.cleanup()
        except OSError:
            pass


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--source', action='store_true')
    mode.add_argument('--executable', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--timeout', type=int, default=90)
    args = parser.parse_args()
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()  # Never accept a stale success report.
    project_root = Path(__file__).resolve().parents[1]
    command = [sys.executable, str(project_root / 'src/app.py')] if args.source else [str(args.executable.resolve())]
    command += ['--smoke-test', '--smoke-output', str(output)]
    env = os.environ.copy()
    env.update(QT_QPA_PLATFORM='offscreen', MPLBACKEND='Agg', PYTHONUTF8='1', PYTHONIOENCODING='utf-8')
    try:
        result = subprocess.run(command, cwd=project_root, env=env, stdout=subprocess.PIPE,
                                stderr=subprocess.STDOUT, text=True, encoding='utf-8', errors='replace', timeout=args.timeout)
        output.with_suffix('.log').write_text(result.stdout or '', encoding='utf-8')
        if result.stdout:
            print(result.stdout)
        if result.returncode != 0:
            raise RuntimeError('Application smoke exited with code ' + str(result.returncode))
        report = json.loads(output.read_text(encoding='utf-8'))
        if not report.get('success'):
            raise RuntimeError('Application smoke reported failure: ' + str(report.get('error')))
        print('Smoke passed: ' + str(output))
        return 0
    except BaseException as exc:
        if isinstance(exc, subprocess.TimeoutExpired):
            captured = exc.stdout or b''
            if isinstance(captured, bytes):
                captured = captured.decode('utf-8', errors='replace')
            output.with_suffix('.log').write_text(captured, encoding='utf-8')
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
